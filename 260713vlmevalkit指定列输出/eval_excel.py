"""Excel export helpers for VLMEvalKit evaluation results.

The writer is intentionally schema-agnostic: detail columns come from the
actual per-sample evaluation data. No dataset-specific field names are hard
coded. ``run.py`` calls :func:`export_eval_results_to_excel` after inference-only
runs and immediately after ``dataset.evaluate`` returns for evaluated runs.
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Mapping, Sequence

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


_INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
_SUPPORTED_TABLE_SUFFIXES = {".xlsx", ".xls", ".csv", ".tsv", ".json", ".jsonl"}


def _json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, default=str)


def _flatten_mapping(data: Mapping[str, Any], prefix: str = "") -> dict[str, Any]:
    flattened: dict[str, Any] = {}
    for key, value in data.items():
        name = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(value, Mapping):
            flattened.update(_flatten_mapping(value, name))
        elif isinstance(value, (list, tuple, set)):
            flattened[name] = _json_text(list(value))
        else:
            flattened[name] = value
    return flattened


def evaluation_result_to_dataframe(eval_results: Any) -> pd.DataFrame | None:
    """Normalize an evaluate() return value for the optional Summary worksheet."""
    if eval_results is None:
        return None
    if isinstance(eval_results, pd.DataFrame):
        return eval_results.copy()
    if isinstance(eval_results, Mapping):
        return pd.DataFrame([_flatten_mapping(eval_results)])
    raise TypeError(
        "Unsupported evaluation result type: "
        f"{type(eval_results).__name__}. Expected None, dict or pandas.DataFrame."
    )


def resolve_export_fields(field_config: Any, dataset_name: str) -> list[str] | None:
    """Resolve optional Detail-sheet fields for one dataset.

    ``None`` means export every column actually present in the detail source.
    A JSON list applies to every dataset. A JSON object may map dataset names
    to lists and use ``*`` as a fallback.
    """
    if field_config is None:
        return None
    if isinstance(field_config, list):
        if not all(isinstance(item, str) for item in field_config):
            raise ValueError("--eval-fields JSON list must contain only strings")
        return field_config
    if isinstance(field_config, Mapping):
        selected = field_config.get(dataset_name, field_config.get("*"))
        if selected is None:
            return None
        if not isinstance(selected, list) or not all(isinstance(item, str) for item in selected):
            raise ValueError(
                f"--eval-fields entry for {dataset_name!r} must be a JSON list of strings"
            )
        return selected
    raise ValueError("--eval-fields must be a JSON list or JSON object")


def _safe_sheet_name(name: str) -> str:
    cleaned = _INVALID_SHEET_CHARS.sub("_", name).strip() or "Sheet"
    return cleaned[:31]


def _read_table(path: Path) -> pd.DataFrame:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xls"}:
        return pd.read_excel(path)
    if suffix == ".csv":
        return pd.read_csv(path)
    if suffix == ".tsv":
        return pd.read_csv(path, sep="\t")
    if suffix == ".jsonl":
        return pd.read_json(path, lines=True)
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, list):
            return pd.DataFrame(data)
        if isinstance(data, Mapping):
            return pd.DataFrame([data])
    raise ValueError(f"Unsupported detail file format: {path}")


def _normalized_token(value: str | None) -> str:
    return re.sub(r"[^a-z0-9]+", "", (value or "").lower())


def find_detail_result_file(
    result_file: str | Path,
    *,
    judge_model: str | None = None,
    output_file: str | Path | None = None,
) -> Path | None:
    """Find the per-sample result produced by evaluate().

    Candidate selection is based on the prediction filename stem, supported
    table formats, optional judge-name affinity, and modification time. It does
    not assume any detail-column names.
    """
    prediction = Path(result_file).expanduser().resolve()
    parent = prediction.parent
    output = Path(output_file).expanduser().resolve() if output_file else None
    stem = prediction.stem
    judge_token = _normalized_token(judge_model)

    candidates: list[tuple[int, float, Path]] = []
    for path in parent.iterdir():
        if not path.is_file() or path.suffix.lower() not in _SUPPORTED_TABLE_SUFFIXES:
            continue
        resolved = path.resolve()
        if resolved == prediction or (output is not None and resolved == output):
            continue

        lower_name = path.name.lower()
        if not path.stem.startswith(stem) or "result" not in lower_name:
            continue

        score = 0
        if "_result" in lower_name:
            score += 20
        if judge_token and judge_token in _normalized_token(path.stem):
            score += 50
        if path.suffix.lower() in {".xlsx", ".csv", ".tsv"}:
            score += 5
        candidates.append((score, path.stat().st_mtime, path))

    if not candidates:
        return None
    candidates.sort(key=lambda item: (item[0], item[1]), reverse=True)
    return candidates[0][2]


def _format_worksheet(worksheet, dataframe: pd.DataFrame) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, column_name in enumerate(dataframe.columns, start=1):
        values = [str(column_name)]
        values.extend("" if pd.isna(value) else str(value) for value in dataframe[column_name].head(200))
        width = min(max(max((len(value) for value in values), default=0) + 2, 10), 60)
        worksheet.column_dimensions[get_column_letter(index)].width = width


def export_eval_results_to_excel(
    eval_results: Any,
    output_file: str | Path,
    *,
    result_file: str | Path | None = None,
    judge_model: str | None = None,
    fields: Sequence[str] | None = None,
) -> tuple[Path, Path | None]:
    """Export per-sample detail and an optional aggregate summary.

    When ``eval_results`` is ``None`` (inference-only mode), the workbook
    contains only ``Detail`` and reads directly from the prediction file.
    When evaluation results are present, the workbook contains ``Detail`` and
    ``Summary`` and prefers the per-sample result produced by evaluate().

    Detail columns are discovered dynamically from the per-sample result file.
    The complete, unfiltered result is first written to a temporary workbook in
    the same directory as the final workbook. The final workbook is then built
    by reading that temporary workbook and optionally applying ``fields``. The
    temporary workbook is removed automatically, so no separate tmp directory
    or manual cleanup is required.

    When ``fields`` is omitted, every available detail column is preserved.
    When no evaluated detail file can be found, the prediction file itself is
    used as a fallback, so original inputs and model predictions are retained.

    Returns:
        ``(workbook_path, detail_source_path)``. In inference-only mode,
        ``detail_source_path`` is the prediction file itself.
    """
    summary_df = evaluation_result_to_dataframe(eval_results)
    inference_only = eval_results is None
    destination = Path(output_file).expanduser().resolve()
    if destination.suffix.lower() != ".xlsx":
        destination = destination.with_suffix(".xlsx")
    destination.parent.mkdir(parents=True, exist_ok=True)

    detail_source: Path | None = None
    detail_df: pd.DataFrame | None = None
    if result_file is not None:
        prediction_file = Path(result_file).expanduser().resolve()
        if inference_only:
            if prediction_file.exists() and prediction_file.suffix.lower() in _SUPPORTED_TABLE_SUFFIXES:
                detail_source = prediction_file
        else:
            detail_source = find_detail_result_file(
                prediction_file,
                judge_model=judge_model,
                output_file=destination,
            )
            if detail_source is None:
                if prediction_file.exists() and prediction_file.suffix.lower() in _SUPPORTED_TABLE_SUFFIXES:
                    detail_source = prediction_file
        if detail_source is not None:
            detail_df = _read_table(detail_source)

    if detail_df is None and summary_df is None:
        raise FileNotFoundError(
            f"No exportable prediction/detail data found for result file: {result_file}"
        )

    # First write the complete, unfiltered data to a temporary workbook in the
    # final result directory. A fixed filename is used so an interrupted earlier
    # run cannot create an ever-growing collection of temporary files.
    temporary = destination.with_name(f"{destination.stem}_full.tmp.xlsx")
    if temporary.exists():
        temporary.unlink()

    try:
        with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
            if detail_df is not None:
                detail_name = _safe_sheet_name("Detail")
                detail_df.to_excel(writer, index=False, sheet_name=detail_name)
                _format_worksheet(writer.book[detail_name], detail_df)

            if summary_df is not None:
                summary_name = _safe_sheet_name("Summary")
                summary_df.to_excel(writer, index=False, sheet_name=summary_name)
                _format_worksheet(writer.book[summary_name], summary_df)

        # Read back the complete export before applying any optional filtering.
        # This keeps the writer schema-agnostic: the final columns come from the
        # actual exported data, not from a hard-coded field list.
        workbook = pd.ExcelFile(temporary)
        final_detail_df: pd.DataFrame | None = None
        if "Detail" in workbook.sheet_names:
            final_detail_df = pd.read_excel(temporary, sheet_name="Detail")
        final_summary_df: pd.DataFrame | None = None
        if "Summary" in workbook.sheet_names:
            final_summary_df = pd.read_excel(temporary, sheet_name="Summary")

        if fields is not None:
            if final_detail_df is None:
                raise ValueError("--eval-fields was provided, but no detail data source was found")
            requested = list(fields)
            missing = [field for field in requested if field not in final_detail_df.columns]
            if missing:
                available = ", ".join(map(str, final_detail_df.columns)) or "<none>"
                raise ValueError(
                    f"Unknown detail field(s): {', '.join(missing)}. Available fields: {available}"
                )
            final_detail_df = final_detail_df.loc[:, requested]

        with pd.ExcelWriter(destination, engine="openpyxl") as writer:
            if final_detail_df is not None:
                detail_name = _safe_sheet_name("Detail")
                final_detail_df.to_excel(writer, index=False, sheet_name=detail_name)
                _format_worksheet(writer.book[detail_name], final_detail_df)

            if final_summary_df is not None:
                summary_name = _safe_sheet_name("Summary")
                final_summary_df.to_excel(writer, index=False, sheet_name=summary_name)
                _format_worksheet(writer.book[summary_name], final_summary_df)
    finally:
        # Automatic cleanup prevents temporary files from accumulating.
        if temporary.exists():
            temporary.unlink()

    return destination, detail_source
