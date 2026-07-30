#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 OpenCompass 的 results + predictions 导出成用于结果分析 / case 分析 / BadCase 分析的 Excel。

相对 export_to_excel.py 的改动：
  1. origin_prompt 是 [{role, prompt}] 列表而不是字符串，按 role 取最后一条 HUMAN，
     再从中剥出真正的 Question 段（去掉 ARB 的 prompt 包装）。
     旧实现 str(list) 会把换行 repr 成字面量 \\n，导致题干列变成一坨转义文本。
  2. 判分同时支持二值 correct 和部分得分 score（prooflike 是 0-10 的 rubric 分，
     没有 correct 字段，旧实现会把 19 题全判成 FALSE）。
  3. details 里的所有字段一律保留：已知字段进固定列，未知字段自动追加成额外列，
     所以 reason / pred_number / lenient_pred 以及以后新增的诊断字段都不会丢。
  4. 额外派生三个用于 BadCase 归因的信号：有无 ANSWER 分隔符、输出字符数、疑似截断。
  5. 多模型 × 多子集在一个工作簿里，明细带「模型 / 子集」列，另出一张 BadCase 表。

v4 修复（重要）：
  6. OC 的 details 把每个字段都包成单元素列表（correct=[False]、rubric_score=[7.0]）。
     旧实现对原始值直接做真假判断，[False] 非空即为真，于是每一道错题都被记成满分：
     2060 题全判"正确"、prooflike 38 题判"无判分"、失分题数整列为 0。
     所有取值一律先过 _unwrap()，且 correct 必须是布尔，否则硬报错不兜底。
  7. 每个子集导出后用逐题得分反推指标并与 results.json 的官方指标对账，
     对不上立即告警。上面那个 bug 之所以能活下来，正是因为没有这道对账。
  8. 总览增加 reason 分类计数列，BadCase 归因不必再回明细做透视。

用法:
    python tools/export_eval_excel.py outputs/arb_full9
    python tools/export_eval_excel.py outputs/arb_full9 --out docs/results/arb_full9.xlsx
    python tools/export_eval_excel.py outputs/arb_full9 --dataset arb_math_prooflike
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise SystemExit('缺少 openpyxl：pip install openpyxl --break-system-packages') from exc

# details 里已经有固定列位的字段，其余的自动追加到末尾
KNOWN_DETAIL_KEYS = {'example_abbr', 'correct', 'score', 'pred', 'answer', 'gold',
                     'origin_prompt', 'predictions', 'references'}
SCORE_KEYS = ('score', 'rubric_score', 'pred_score')
METRIC_KEYS = ('accuracy', 'score', 'rubric_score')

BASE_COLS = ['模型', '子集', '题号', '判定', '得分', '满分',
             '模型最终答案', '标准答案', '有ANSWER分隔符', '输出字符数', '疑似截断',
             '题干', '完整输出']

OVERVIEW_BASE = ['模型', '子集', '题数', '官方指标', '逐题重算', '对账',
                 '单题权重', '判分方式', '失分题数',
                 '空输出', '缺ANSWER分隔符', '疑似截断']

# ARB 四套 prompt 的公共包装，用来把真正的题面抠出来
_Q_START = re.compile(r'Question:\s*', re.S)
_Q_END = re.compile(r'\n\s*Now it is time to', re.S)


def resolve_work_dir(path):
    """OC 用 -w 时会在下面再套一层时间戳目录。传 outputs/arb_full9 或
    outputs/arb_full9/20260729_051215 都要能工作。"""
    from pathlib import Path
    p = Path(path)
    if not p.is_dir():
        raise SystemExit(f'{p} 不存在')
    if (p / 'results').is_dir():
        return p
    cands = sorted(c for c in p.iterdir() if c.is_dir() and (c / 'results').is_dir())
    if not cands:
        raise SystemExit(f'{p} 下面找不到 results/ —— 确认这是 OC 的 work_dir')
    if len(cands) > 1:
        print(f'[!] {p} 下有 {len(cands)} 个运行目录，用最新的 {cands[-1].name}')
        print(f'    其余: {[c.name for c in cands[:-1]]}')
    return cands[-1]


def _read(p: Path) -> Any:
    return json.loads(p.read_text(encoding='utf-8'))


def _unwrap(v: Any) -> Any:
    """OC 的 details 把每个字段都包成单元素列表：correct=[False]、rubric_score=[7.0]。

    不拆包直接做真假判断的话，[False] 因为列表非空而为真，所有错题都会被静默
    记成满分。凡是要参与判分的取值，一律先过这里。
    """
    while isinstance(v, list) and len(v) == 1:
        v = v[0]
    return v


def question_text(origin_prompt: Any) -> str:
    """origin_prompt 可能是 [{role, prompt}] 列表，也可能是裸字符串。"""
    if isinstance(origin_prompt, list):
        humans = [t.get('prompt', '') for t in origin_prompt
                  if isinstance(t, dict) and str(t.get('role', '')).upper() == 'HUMAN']
        text = humans[-1] if humans else ''
        if not text:
            text = '\n'.join(str(t.get('prompt', '')) for t in origin_prompt
                             if isinstance(t, dict))
    else:
        text = str(origin_prompt or '')
    m = _Q_START.search(text)
    if m:
        text = text[m.end():]
    m = _Q_END.search(text)
    if m:
        text = text[:m.start()]
    return text.strip()


def has_delimiter(pred_text: str) -> bool:
    return bool(re.search(r'ANSWER\s*:', pred_text or ''))


def looks_truncated(pred_text: str) -> bool:
    """没有 ANSWER 分隔符、且结尾不是句子终止符 —— 大概率是 max_out_len 截断。"""
    t = (pred_text or '').rstrip()
    if not t:
        return False
    if has_delimiter(t):
        return False
    return t[-1] not in '.。!！?？)）]}$'


def index_of(abbr: Any, fallback: int) -> str:
    m = re.search(r'(\d+)$', str(abbr or ''))
    return m.group(1) if m else str(fallback)


def flatten(v: Any) -> str:
    if isinstance(v, list):
        if len(v) == 1:
            return flatten(v[0])
        return ' | '.join(flatten(x) for x in v)
    if v is None:
        return ''
    return str(v)


def load_pair(work_dir: Path, model: str, dataset: str):
    res = _read(work_dir / 'results' / model / f'{dataset}.json')
    pred_path = work_dir / 'predictions' / model / f'{dataset}.json'
    preds = _read(pred_path) if pred_path.is_file() else {}
    if isinstance(preds, list):
        preds = {str(i): v for i, v in enumerate(preds)}
    details = res.get('details')
    if isinstance(details, dict):
        details = list(details.values())
    details = [d for d in (details or []) if isinstance(d, dict)]
    metric = next((float(res[k]) for k in METRIC_KEYS
                   if isinstance(res.get(k), (int, float))), None)
    return metric, details, preds


def build_rows(work_dir: Path, only_dataset: str | None):
    rows, extra_keys, overview = [], [], []
    reason_keys: list[str] = []
    work_dir = resolve_work_dir(work_dir)
    res_root = work_dir / 'results'

    for model_dir in sorted(p for p in res_root.iterdir() if p.is_dir()):
        model = model_dir.name
        for f in sorted(model_dir.glob('*.json')):
            dataset = f.stem
            if only_dataset and dataset != only_dataset:
                continue
            metric, details, preds = load_pair(work_dir, model, dataset)

            scores = [_unwrap(next((d[k] for k in SCORE_KEYS if k in d), None))
                      for d in details]
            partial = any(isinstance(s, (int, float)) and not isinstance(s, bool)
                          and s not in (0, 1) for s in scores)
            full = 10.0 if partial else 1.0

            n_bad = n_empty = n_nodelim = n_trunc = 0
            reason_count: dict[str, int] = {}
            earned = 0.0
            for i, d in enumerate(details):
                idx = index_of(d.get('example_abbr'), i)
                rec = preds.get(idx) or preds.get(str(i)) or {}
                raw = flatten(rec.get('prediction', ''))

                raw_score = _unwrap(next((d[k] for k in SCORE_KEYS if k in d), None))
                if isinstance(raw_score, bool):
                    raw_score = 1.0 if raw_score else 0.0
                if raw_score is None and 'correct' in d:
                    c = _unwrap(d['correct'])
                    if not isinstance(c, bool):
                        raise SystemExit(
                            f'{model}/{dataset} 第 {idx} 题 correct 非布尔值: {c!r} '
                            f'—— 判分路径不做兜底，请先确认 evaluator 的输出格式')
                    raw_score = 1.0 if c else 0.0
                score = float(raw_score) if isinstance(raw_score, (int, float)) else None

                if score is None:
                    verdict = '无判分'
                elif score >= full:
                    verdict = '正确'
                elif score <= 0:
                    verdict = '错误'
                else:
                    verdict = f'部分得分 {score:g}/{full:g}'

                delim, trunc = has_delimiter(raw), looks_truncated(raw)
                if score is not None:
                    earned += score
                    if score < full:
                        n_bad += 1
                if not raw.strip():
                    n_empty += 1
                if not delim:
                    n_nodelim += 1
                if trunc:
                    n_trunc += 1

                reason = flatten(d.get('reason', '')).strip()
                if reason:
                    reason_count[reason] = reason_count.get(reason, 0) + 1
                    if reason not in reason_keys:
                        reason_keys.append(reason)

                row = {
                    '模型': model, '子集': dataset, '题号': idx,
                    '判定': verdict,
                    '得分': score if score is not None else '',
                    '满分': full,
                    '模型最终答案': flatten(d.get('pred', '')),
                    '标准答案': flatten(d.get('answer', d.get('gold', ''))),
                    '有ANSWER分隔符': delim,
                    '输出字符数': len(raw),
                    '疑似截断': trunc,
                    '题干': question_text(rec.get('origin_prompt', '')),
                    '完整输出': raw,
                }
                for k, v in d.items():
                    if k in KNOWN_DETAIL_KEYS:
                        continue
                    if k not in extra_keys:
                        extra_keys.append(k)
                    row[k] = flatten(v)
                rows.append(row)

            # 对账：逐题得分反推指标，必须与 results.json 的官方指标一致。
            # 判分路径一旦出错（例如未拆包导致全判对），这里会立刻暴露。
            recomputed = earned / (len(details) * full) * 100 if details else None
            if metric is None or recomputed is None:
                verdict_check = '无官方指标'
            elif abs(recomputed - metric) <= 0.01:
                verdict_check = 'OK'
            else:
                verdict_check = f'不一致 差{recomputed - metric:+.2f}'
                print(f'  [!] {model}/{dataset} 逐题重算 {recomputed:.2f} '
                      f'≠ 官方指标 {metric:.2f} —— 判分路径有问题，不要直接使用本表')

            rec_ov = {
                '模型': model, '子集': dataset, '题数': len(details),
                '官方指标': metric if metric is not None else '',
                '逐题重算': round(recomputed, 4) if recomputed is not None else '',
                '对账': verdict_check,
                '单题权重': round(100 / len(details), 2) if details else '',
                '判分方式': f'部分得分 0-{full:g}' if partial else '二值',
                '失分题数': n_bad,
                '空输出': n_empty, '缺ANSWER分隔符': n_nodelim, '疑似截断': n_trunc,
            }
            rec_ov.update(reason_count)
            overview.append(rec_ov)

    # 各子集出现的 reason 不同，统一补齐成同一组列，缺的填 0
    for o in overview:
        for k in reason_keys:
            o.setdefault(k, 0)
    overview = [{c: o[c] for c in OVERVIEW_BASE + reason_keys} for o in overview]
    return rows, extra_keys, overview


HEAD_FILL = PatternFill('solid', fgColor='1F4E78')
HEAD_FONT = Font(bold=True, color='FFFFFF')
OK_FILL = PatternFill('solid', fgColor='E2F0D9')
BAD_FILL = PatternFill('solid', fgColor='FCE4D6')
WARN_FILL = PatternFill('solid', fgColor='FFF2CC')

WIDTHS = {'模型': 18, '子集': 22, '题号': 8, '判定': 15, '得分': 8, '满分': 8,
          '模型最终答案': 32, '标准答案': 28, '有ANSWER分隔符': 15,
          '输出字符数': 12, '疑似截断': 10, '题干': 60, '完整输出': 80,
          'reason': 45, 'pred_number': 14, 'lenient_pred': 22,
          '题数': 8, '官方指标': 12, '逐题重算': 12, '对账': 16,
          '单题权重': 12, '判分方式': 16,
          '失分题数': 12, '空输出': 10, '缺ANSWER分隔符': 16}


def write_sheet(ws, cols, rows, freeze='A2'):
    ws.append(cols)
    for r in rows:
        ws.append([r.get(c, '') for c in cols])
    for cell in ws[1]:
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    if '判定' in cols:
        c = cols.index('判定') + 1
        for r in range(2, ws.max_row + 1):
            v = str(ws.cell(r, c).value)
            ws.cell(r, c).fill = OK_FILL if v == '正确' else (
                BAD_FILL if v == '错误' else WARN_FILL)
    if '对账' in cols:
        c = cols.index('对账') + 1
        for r in range(2, ws.max_row + 1):
            v = str(ws.cell(r, c).value)
            ws.cell(r, c).fill = OK_FILL if v == 'OK' else BAD_FILL
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(c, 18)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('work_dir')
    ap.add_argument('--out', default=None)
    ap.add_argument('--dataset', default=None)
    args = ap.parse_args()

    wd = resolve_work_dir(args.work_dir)
    rows, extra_keys, overview = build_rows(wd, args.dataset)
    if not rows:
        raise SystemExit('没读到任何逐题数据，检查 work_dir / --dataset')

    cols = BASE_COLS[:]
    for k in ('reason', 'pred_number', 'lenient_pred'):
        if k in extra_keys:
            cols.insert(cols.index('标准答案') + 1, k)
    cols += [k for k in extra_keys if k not in cols]

    bad = [r for r in rows if r['判定'] != '正确']

    wb = Workbook()
    write_sheet(wb.active, list(overview[0].keys()), overview)
    wb.active.title = '总览'
    write_sheet(wb.create_sheet('明细'), cols, rows)
    write_sheet(wb.create_sheet('BadCase'), cols, bad)
    wb.active = 0

    out = Path(args.out) if args.out else wd / f'{wd.name}_analysis.xlsx'
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)

    n_mismatch = sum(1 for o in overview if o['对账'] != 'OK')
    print(f'写出 {out}')
    print(f'  总览 {len(overview)} 行 / 明细 {len(rows)} 行 / BadCase {len(bad)} 行')
    print(f'  保留的诊断字段: {extra_keys or "无"}')
    for o in overview:
        print(f"  {o['模型']:<18}{o['子集']:<24}n={o['题数']:<5}"
              f"失分={o['失分题数']:<4}空={o['空输出']:<3}"
              f"缺分隔符={o['缺ANSWER分隔符']:<3}疑似截断={o['疑似截断']:<3}"
              f"对账={o['对账']}")
    if n_mismatch:
        raise SystemExit(f'\n[!] {n_mismatch} 个子集对账失败，本表不可用于报告')
    print('\n对账全部通过：逐题得分与官方指标一致。')


if __name__ == '__main__':
    main()
